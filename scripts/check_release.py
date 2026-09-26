"""Deterministic local integrity checks for an honest item-3 working release."""
import json
import subprocess
from pathlib import Path
import pandas as pd
from docx import Document
from scipy.stats import binomtest

root=Path(__file__).resolve().parents[1]
tools=pd.read_csv(root/'results/tools_ledger.csv')
used=tools[tools.counts_for_gate]
assert len(used)==40 and used.tool.nunique()==40
sets=pd.read_csv(root/'results/datasets_ledger.csv')
primary=sets[sets.role=='primary']
assert len(primary)==168 and primary.accession.nunique()==168
size=json.loads((root/'results/blocked_size.json').read_text())
assert (size['n_eligible_donors'],size['n_positive_donors'])==(12,9)
assert abs(size['sign_p_one_sided']-binomtest(9,12,.5,alternative='greater').pvalue)<1e-12
assert size['H1']['verdict']=='FAIL'
zero=json.loads((root/'results/zero_fsk_control.json').read_text())
assert (zero['eligible_donors'],zero['positive_paired_delta_donors'],zero['verdict'])==(11,3,'FAIL')
assert abs(zero['exact_one_sided_sign_p']-binomtest(3,11,.5,alternative='greater').pvalue)<1e-12
seg=json.loads((root/'results/trainonly_seg/sealed_eval.json').read_text())
assert len(seg['per_image'])==12 and abs(seg['mean_author']-.738672)<1e-6
assert seg['mean_author']<seg['published_comparison']==.76
selection=json.loads((root/'results/trainonly_seg/val_selection.json').read_text())
assert selection['eval_used'] is False and selection['selected_params']==[.4,.5,.3,40]
assert (root/'results/trainonly_seg/best.pt').is_file()
paper=Document(root/'paper/mega27-03-virtual-organoid-paper.docx')
pdf=root/'paper/mega27-03-virtual-organoid-paper.pdf'
assert pdf.is_file()
info=subprocess.check_output(['pdfinfo',str(pdf)],text=True)
pages=int(next(line.split(':',1)[1].strip() for line in info.splitlines() if line.startswith('Pages:')))
assert pages>=50
fonts=subprocess.check_output(['pdffonts',str(pdf)],text=True)
assert 'TimesNewRomanPSMT' in fonts and 'TimesNewRomanPS-BoldMT' in fonts
assert all(' yes ' in row for row in fonts.splitlines()[2:] if 'TimesNewRoman' in row)
assert len(paper.tables)==51
text='\n'.join(x.text for x in paper.paragraphs)
assert 'a paired no-forskolin specificity test both failed' in text
assert '3/11 donor-median paired deltas are positive' in text
assert 'NOT an established new biological discovery' in text
print('PASS: 40 unique tools, 168 unique primary entries, both negative size checks, '
      f'sealed segmentation below .76, frozen selection, checkpoint, {pages}-page Times New Roman PDF, 51 tables and explicit caveats')
